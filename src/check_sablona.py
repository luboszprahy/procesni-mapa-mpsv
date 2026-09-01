# -*- coding: utf-8 -*-
"""Brána nad vygenerovanou šablonou pro hromadný import aktivit.

Šablona je jediný artefakt projektu, který se dostane do ruky správci a zpátky
se čte strojem. Vada v ní se pozná až tím, že import „nic nevidí" nebo že
založí aktivity s posunutými sloupci — obojí vypadá jako chyba flow, ne jako
chyba sešitu. Proto se kontroluje offline, nad tím, co se opravdu vydává:

1. tvar sešitu — právě dva listy, data v „Aktivity“, pokyny mimo tabulku;
2. formátovaná Tabulka, ne volná mřížka — bez ní konektor Excel Online
   `List rows present in a table` neuvidí nic a vypadá to jako oprávnění;
3. hlavička = sloupce schématu minus systémové, ve stejném pořadí;
4. systémové sloupce v šabloně NEJSOU — kód vyplněný ručně je vada, ne vstup;
5. šablona je prázdná — vydaný sešit nesmí vézt zbytková data;
6. Choice sloupce mají rozbalovátko s volbami ze schématu;
7. list Pokyny popisuje přesně sloupce tabulky a jeho tvrzení sedí na schéma.

Spouštět z kořene projektu:
    python src/check_sablona.py
"""

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

sys.path.insert(0, "src")
from make_sablona import (LIST, LIST_POKYNY, NAPOVEDA, SCHEMA,  # noqa: E402
                          TABULKA, VYSTUP, nacti_schema, omezeni,
                          sloupce_sablony)

# Záměrně vlastní seznam, ne SYSTEMOVE z generátoru: kdyby z něj sloupec
# vypadl, generátor i brána by se shodly na tom, že do šablony patří, a nikdo
# by nezakřičel. Tohle je druhé, nezávislé vyslovení téhož pravidla.
ZAKAZANE = {
    "Title": "identifikační kód přiděluje appka",
    "nazev_kratky": "odvozuje se z názvu",
    "datum_aktualizace": "razítko importu",
    "puvodni_kod": "stopa po přesunu pod jiného rodiče",
}

chyby = []
kontrol = 0


def overit(podminka, popis):
    global kontrol
    kontrol += 1
    if not podminka:
        chyby.append(popis)


def zkontroluj_tvar_sesitu(sesit):
    overit(sesit.sheetnames == [LIST, LIST_POKYNY],
           f"listy sešitu jsou {sesit.sheetnames}, čekám {[LIST, LIST_POKYNY]}")
    overit(not sesit[LIST_POKYNY].tables if LIST_POKYNY in sesit.sheetnames else True,
           f"list {LIST_POKYNY} obsahuje Tabulku — konektor by mohl sáhnout na ni")


def zkontroluj_tabulku(list_dat, pocet_sloupcu):
    """Vrátí hlavičku Tabulky, nebo None, když Tabulka není použitelná."""
    overit(len(list_dat.tables) == 1,
           f"na listu {LIST} je {len(list_dat.tables)} Tabulek, čekám právě jednu")
    if len(list_dat.tables) != 1:
        return None

    jmeno = next(iter(list_dat.tables))
    tabulka = list_dat.tables[jmeno]
    overit(jmeno == TABULKA, f"Tabulka se jmenuje {jmeno!r}, ImportFlow čeká {TABULKA!r}")
    overit(tabulka.headerRowCount == 1,
           f"Tabulka má headerRowCount {tabulka.headerRowCount}, čekám 1")

    zacatek, konec = tabulka.ref.split(":")
    overit(zacatek == "A1", f"Tabulka začíná na {zacatek}, čekám A1")
    ocekavany_konec = f"{get_column_letter(pocet_sloupcu)}2"
    overit(konec == ocekavany_konec,
           f"Tabulka končí na {konec}, čekám {ocekavany_konec} "
           f"(hlavička + jeden prázdný řádek)")

    # Konektor nečte buňky hlavičky, ale jména sloupců zapsaná v definici
    # Tabulky. Rozejít se můžou — a rozejdou-li se, import čte jinou sadu klíčů,
    # než jakou má správce před očima.
    hlavicka = [b.value for b in list_dat[1][:pocet_sloupcu]]
    v_definici = [s.name for s in tabulka.tableColumns]
    overit(v_definici == hlavicka,
           f"jména sloupců v definici Tabulky {v_definici} "
           f"neodpovídají hlavičce {hlavicka}")
    return hlavicka


def zkontroluj_hlavicku(hlavicka, sloupce, schema):
    for index, sloupec in enumerate(sloupce):
        skutecny = hlavicka[index] if index < len(hlavicka) else None
        overit(skutecny == sloupec["display"],
               f"sloupec {index + 1} má hlavičku {skutecny!r}, "
               f"schéma říká {sloupec['display']!r}")

    aktivity = next(l for l in schema["lists"] if l["name"] == LIST)
    podle_nazvu = {c["name"]: c for c in aktivity["columns"]}
    for systemovy, duvod in ZAKAZANE.items():
        display = podle_nazvu.get(systemovy, {}).get("display")
        overit(display not in hlavicka,
               f"v šabloně je systémový sloupec {display!r} ({duvod})")


def zkontroluj_prazdnost(list_dat, pocet_sloupcu):
    vyplnene = [
        f"{get_column_letter(bunka.column)}{bunka.row}"
        for radek in list_dat.iter_rows(min_row=2, max_col=pocet_sloupcu)
        for bunka in radek
        if bunka.value not in (None, "")
    ]
    overit(not vyplnene, f"vydaná šablona veze data v buňkách {vyplnene[:5]}")


def zkontroluj_rozbalovatka(list_dat, sloupce):
    rozsahy = {}
    for kontrola in list_dat.data_validations.dataValidation:
        for oblast in kontrola.sqref.ranges:
            rozsahy[str(oblast)] = kontrola

    for index, sloupec in enumerate(sloupce, start=1):
        pismeno = get_column_letter(index)
        moje = [k for oblast, k in rozsahy.items() if oblast.startswith(f"{pismeno}2:")]

        if sloupec["type"] != "Choice":
            overit(not moje,
                   f"sloupec {sloupec['display']!r} není Choice, přesto má rozbalovátko")
            continue

        overit(len(moje) == 1,
               f"sloupec {sloupec['display']!r} má {len(moje)} rozbalovátek, čekám jedno")
        if len(moje) != 1:
            continue
        kontrola = moje[0]
        overit(kontrola.type == "list",
               f"rozbalovátko u {sloupec['display']!r} je typu {kontrola.type!r}")
        overit(bool(kontrola.showErrorMessage),
               f"rozbalovátko u {sloupec['display']!r} nehlásí chybu — "
               f"překlep by prošel až k importu")
        volby = (kontrola.formula1 or "").strip('"').split(",")
        overit(volby == sloupec["choices"],
               f"rozbalovátko u {sloupec['display']!r} nabízí {volby}, "
               f"schéma má {sloupec['choices']}")


def zkontroluj_pokyny(list_pokyny, sloupce):
    radky = list(list_pokyny.iter_rows(values_only=True))
    zacatek = next((i + 1 for i, r in enumerate(radky) if r and r[0] == "Sloupec"), None)
    overit(zacatek is not None, f"na listu {LIST_POKYNY} chybí hlavička tabulky pokynů")
    if zacatek is None:
        return

    popsane = [r for r in radky[zacatek:] if r and r[0]]
    overit(len(popsane) == len(sloupce),
           f"pokyny popisují {len(popsane)} sloupců, šablona jich má {len(sloupce)}")

    for index, sloupec in enumerate(sloupce):
        if index >= len(popsane):
            break
        display, povinny, omezeni_text, napoveda = popsane[index][:4]
        overit(display == sloupec["display"],
               f"pokyn {index + 1} popisuje {display!r}, v šabloně je "
               f"{sloupec['display']!r}")
        overit(povinny == ("ano" if sloupec.get("required") else "ne"),
               f"pokyn u {sloupec['display']!r} tvrdí povinnost {povinny!r}, "
               f"schéma říká {'ano' if sloupec.get('required') else 'ne'}")
        overit(omezeni_text == omezeni(sloupec),
               f"pokyn u {sloupec['display']!r} tvrdí omezení {omezeni_text!r}, "
               f"ze schématu plyne {omezeni(sloupec)!r}")
        overit(napoveda == NAPOVEDA.get(sloupec["name"]),
               f"nápověda u {sloupec['display']!r} neodpovídá NAPOVEDA")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--sablona", default=str(VYSTUP))
    parser.add_argument("--schema", default=str(SCHEMA))
    argumenty = parser.parse_args()

    cesta = Path(argumenty.sablona)
    if not cesta.exists():
        raise SystemExit(f"CHYBA: šablona {cesta} neexistuje — spusť make_sablona.py")

    schema = nacti_schema(argumenty.schema)
    sloupce = sloupce_sablony(schema)
    sesit = load_workbook(cesta)

    zkontroluj_tvar_sesitu(sesit)
    if LIST not in sesit.sheetnames:
        vypis()
        return 1

    list_dat = sesit[LIST]
    hlavicka = zkontroluj_tabulku(list_dat, len(sloupce))
    if hlavicka is not None:
        zkontroluj_hlavicku(hlavicka, sloupce, schema)
    zkontroluj_prazdnost(list_dat, len(sloupce))
    zkontroluj_rozbalovatka(list_dat, sloupce)
    if LIST_POKYNY in sesit.sheetnames:
        zkontroluj_pokyny(sesit[LIST_POKYNY], sloupce)

    vypis()
    return 1 if chyby else 0


def vypis():
    print(f"kontrol: {kontrol}")
    for chyba in chyby:
        print(f"CHYBA: {chyba}")
    print("NEPROŠLO" if chyby else "OK — šablona odpovídá schématu")


if __name__ == "__main__":
    sys.exit(main())
