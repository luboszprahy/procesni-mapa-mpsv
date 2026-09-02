# -*- coding: utf-8 -*-
"""Mutační test brány check_sablona.py.

Zelená brána nad souborem, který sama nekontroluje pořádně, nedokazuje nic.
Tenhle skript zavede po jedné chybě — jednou do schématu, jednou do sešitu —
a ověří, že ji brána shodí. Vybrané jsou ty, které by v provozu prošly bez
jediného varování: volná mřížka místo Tabulky (konektor „nic nevidí"),
posunuté pořadí sloupců (import zapíše útvar do předpisu), rozejití pokynů
se schématem (správce vyplní, co se pak odmítne).

Spouštět z kořene projektu:
    python src/mutace_sablona.py
"""

import argparse
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo

sys.path.insert(0, "src")
from make_sablona import LIST, LIST_POKYNY, SCHEMA, TABULKA, VYSTUP  # noqa: E402


def aktivity(schema):
    return next(l for l in schema["lists"] if l["name"] == LIST)


# ---------- mutace schématu (sešit zůstává, jak byl vydán) ----------

def pribyl_sloupec(schema):
    """Sloupec přidaný do schématu, ale nedopsaný do šablony."""
    aktivity(schema)["columns"].append(
        {"name": "priorita", "display": "Priorita", "type": "Text", "maxlen": 20})


def jiny_display(schema):
    """Přejmenovaný sloupec: hlavička v sešitě zůstane stará."""
    next(c for c in aktivity(schema)["columns"] if c["name"] == "vykonava")["display"] = \
        "Vykonává oddělení"


def jine_volby(schema):
    """Do Choice přibyla volba — rozbalovátko ji nenabídne."""
    next(c for c in aktivity(schema)["columns"] if c["name"] == "stav")["choices"] = \
        ["pracovní", "schváleno", "zrušeno"]


def jina_povinnost(schema):
    """Sloupec přestal být povinný, pokyny to pořád tvrdí."""
    next(c for c in aktivity(schema)["columns"] if c["name"] == "dilci_proces_kod") \
        .pop("required")


def jina_delka(schema):
    """Zkrácený sloupec — pokyny slibují víc znaků, než list unese."""
    next(c for c in aktivity(schema)["columns"] if c["name"] == "sekce")["maxlen"] = 10


def ubyl_sloupec(schema):
    """Sloupec ze schématu zmizel, v šabloně zůstal."""
    sloupce = aktivity(schema)["columns"]
    sloupce.remove(next(c for c in sloupce if c["name"] == "sekce"))


# ---------- mutace sešitu (schéma zůstává, jak je) ----------

def bez_tabulky(sesit):
    """Volná mřížka: konektor neuvidí nic a vypadá to jako chyba oprávnění."""
    del sesit[LIST].tables[TABULKA]


def prejmenovana_tabulka(sesit):
    """Jiné jméno Tabulky — ImportFlow sáhne na neexistující."""
    list_dat = sesit[LIST]
    puvodni = list_dat.tables[TABULKA].ref
    del list_dat.tables[TABULKA]
    nova = Table(displayName="Tabulka1", ref=puvodni)
    nova.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    list_dat.add_table(nova)


def prohozene_sloupce(sesit):
    """Nejtišší vada ze všech: import zapíše útvar do předpisu."""
    list_dat = sesit[LIST]
    list_dat["C1"].value, list_dat["D1"].value = \
        list_dat["D1"].value, list_dat["C1"].value
    for sloupec in list_dat.tables[TABULKA].tableColumns:
        if sloupec.name == "Vykonává útvar":
            sloupec.name = "Spolupracuje"
        elif sloupec.name == "Spolupracuje":
            sloupec.name = "Vykonává útvar"


def systemovy_sloupec(sesit):
    """Sloupec pro kód: správce ho vyplní a import by respektoval cizí kódy."""
    list_dat = sesit[LIST]
    list_dat["A1"].value = "Kód"
    list_dat.tables[TABULKA].tableColumns[0].name = "Kód"


def zbyl_odvozeny_sloupec(sesit):
    """Tvar, jaký by vydal generátor, kterému ze SYSTEMOVE vypadl sloupec:
    šablona žádá odvozený údaj, který si systém drží sám."""
    list_dat = sesit[LIST]
    tabulka = list_dat.tables[TABULKA]
    list_dat["I1"].value = "Vzniklo z kódu"
    tabulka.tableColumns.append(TableColumn(id=9, name="Vzniklo z kódu"))
    tabulka.ref = tabulka.ref.replace("H2", "I2")


def zbytkova_data(sesit):
    """Vydaná šablona s daty — správce je přehlédne a naimportuje."""
    sesit[LIST]["A2"].value = "zkušební aktivita"


def bez_rozbalovatka(sesit):
    """Bez rozbalovátka se do Stav dostane cokoli a import to odmítne."""
    sesit[LIST].data_validations.dataValidation = []


def rozbalovatko_bez_hlasky(sesit):
    """Rozbalovátko jako ozdoba: nabídne volby, ale překlep pustí dál."""
    for kontrola in sesit[LIST].data_validations.dataValidation:
        kontrola.showErrorMessage = False


def kratsi_pokyny(sesit):
    """Pokyny zapomenou na poslední sloupec."""
    list_pokyny = sesit[LIST_POKYNY]
    list_pokyny.delete_rows(list_pokyny.max_row)


def kaskada_na_cizi_rozsah(sesit):
    """Pojmenovaný rozsah jednoho procesu ukazuje na dílčí procesy jiného.

    Nejzákeřnější z mutací: nabídka se otevře, něco v ní je, a správce
    z ní vybere kód, který pod jeho proces nepatří. Import ho přijme —
    ten dílčí proces existuje — a aktivita skončí jinde, než měla.
    """
    from openpyxl.workbook.defined_name import DefinedName
    jmeno = "P_01_01"
    del sesit.defined_names[jmeno]
    sesit.defined_names.add(DefinedName(jmeno, attr_text="Ciselniky!$D$40:$D$45"))


def rozsah_do_zasoby(sesit):
    """Rozsah útvarů natažený na strop: nabídka se naplní prázdnými řádky."""
    from openpyxl.workbook.defined_name import DefinedName
    del sesit.defined_names["Cis_Utvary"]
    sesit.defined_names.add(
        DefinedName("Cis_Utvary", attr_text="Ciselniky!$H$2:$H$201"))


def dilci_proces_bez_kaskady(sesit):
    """Dílčí proces dostane plochý seznam všech 250 kódů místo kaskády —
    rozbalovátko sice je, ale roluje se v něm k neupotřebení."""
    for kontrola in sesit[LIST].data_validations.dataValidation:
        if str(next(iter(kontrola.sqref.ranges))).startswith("C2:"):
            kontrola.formula1 = "=Cis_Procesy"


def ciselnik_bez_prazdna(sesit):
    """Rozbalovátko nepustí prázdno — aktivitu bez zařazení pak nejde nahrát."""
    for kontrola in sesit[LIST].data_validations.dataValidation:
        kontrola.allowBlank = False


MUTACE = [
    ("do schématu přibyl sloupec", "schema", pribyl_sloupec),
    ("sloupec se ve schématu přejmenoval", "schema", jiny_display),
    ("Choice má ve schématu novou volbu", "schema", jine_volby),
    ("sloupec přestal být povinný", "schema", jina_povinnost),
    ("sloupec se ve schématu zkrátil", "schema", jina_delka),
    ("sloupec ze schématu zmizel", "schema", ubyl_sloupec),
    ("volná mřížka místo Tabulky", "sesit", bez_tabulky),
    ("přejmenovaná Tabulka", "sesit", prejmenovana_tabulka),
    ("prohozené pořadí dvou sloupců", "sesit", prohozene_sloupce),
    ("v šabloně je sloupec pro kód", "sesit", systemovy_sloupec),
    ("v šabloně zbyl odvozený sloupec", "sesit", zbyl_odvozeny_sloupec),
    ("vydaná šablona veze data", "sesit", zbytkova_data),
    ("Stav bez rozbalovátka", "sesit", bez_rozbalovatka),
    ("rozbalovátko bez chybové hlášky", "sesit", rozbalovatko_bez_hlasky),
    ("pokyny zapomněly na sloupec", "sesit", kratsi_pokyny),
    ("kaskáda ukazuje na cizí proces", "sesit", kaskada_na_cizi_rozsah),
    ("rozsah číselníku natažený do zásoby", "sesit", rozsah_do_zasoby),
    ("dílčí proces bez kaskády", "sesit", dilci_proces_bez_kaskady),
    ("rozbalovátko nepustí prázdno", "sesit", ciselnik_bez_prazdna),
]


def priprav(docasny, sablona, schema, kde, zmena):
    """Vrátí dvojici (cesta k šabloně, cesta ke schématu) po zavedení mutace."""
    kopie_sablony = Path(docasny) / "sablona.xlsx"
    kopie_schematu = Path(docasny) / "schema.json"
    shutil.copy(sablona, kopie_sablony)
    shutil.copy(schema, kopie_schematu)

    if kde == "schema":
        obsah = json.loads(kopie_schematu.read_text(encoding="utf-8"))
        pred = json.dumps(obsah, ensure_ascii=False, sort_keys=True)
        zmena(obsah)
        if json.dumps(obsah, ensure_ascii=False, sort_keys=True) == pred:
            raise SystemExit(f"CHYBA: mutace {zmena.__name__} se do schématu netrefila")
        kopie_schematu.write_text(
            json.dumps(obsah, ensure_ascii=False, indent=2), encoding="utf-8")
    else:
        sesit = load_workbook(kopie_sablony)
        zmena(sesit)
        sesit.save(kopie_sablony)

    return kopie_sablony, kopie_schematu


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--sablona", default=str(VYSTUP))
    parser.add_argument("--schema", default=str(SCHEMA))
    argumenty = parser.parse_args()

    hlidane = 0
    for popis, kde, zmena in MUTACE:
        with tempfile.TemporaryDirectory() as docasny:
            sablona, schema = priprav(
                docasny, argumenty.sablona, argumenty.schema, kde, zmena)
            beh = subprocess.run(
                [sys.executable, "src/check_sablona.py",
                 "--sablona", str(sablona), "--schema", str(schema)],
                capture_output=True, text=True, encoding="utf-8", errors="replace")

        if beh.returncode == 0:
            print(f"NEODHALENO: {popis} — brána mutaci propustila")
        else:
            hlidane += 1
            # První hláška, ne jen návratový kód: mutace chycená nesouvisející
            # kontrolou vypadá stejně zeleně a přitom nedokazuje nic.
            prvni = next((r for r in beh.stdout.splitlines()
                          if r.startswith("CHYBA:")), beh.stderr.strip()[:150])
            print(f"chycena: {popis}")
            print(f"          -> {prvni[:150]}")

    print(f"\nchycených mutací: {hlidane}/{len(MUTACE)}")
    return 0 if hlidane == len(MUTACE) else 1


if __name__ == "__main__":
    sys.exit(main())
